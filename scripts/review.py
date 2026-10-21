#!/usr/bin/env python3
"""La review hebdo client : un fichier Excel par compte, en aller-retour avec Metricool.

    python3 scripts/review.py generer   [--semaine 2026-08-10] [--client le-loft] [--dry-run]
    python3 scripts/review.py appliquer [--client le-loft] [--dry-run]

Le circuit complet est décrit dans .claude/skills/cm-metricool/references/review-hebdo.md.
En résumé : `generer` lit les brouillons de la semaine dans Metricool et écrit un fichier
Excel par client dans reviews/a-envoyer/ ; la session Claude du lundi le dépose dans le
Drive du client (converti en Google Sheet) et note son identifiant dans reviews/config.json ;
le client coche ✅ / ✏️ / ⛔ ligne par ligne, corrige les légendes ; `appliquer`, plusieurs
fois par jour, relit le fichier par son export public, pousse les modifications dans
Metricool et passe les lignes ✅ en publication automatique.

Même partage des rôles que le reste du dépôt : ce script tourne dans GitHub Actions — la
politique réseau des sessions Claude bloque app.metricool.com et les exports Drive — et la
session fait la moitié authentifiée (dépôt du fichier, récap). Aucune clé Google, comme
drive_sync.py : la lecture des retours passe par l'export public du fichier partagé.

Garde-fous, dans l'ordre où ils ont été décidés :
  - jamais de suppression automatique : un ⛔ reste en brouillon, Maxime tranche ;
  - un post déjà validé dans Metricool (autoPublish) n'est JAMAIS retouché ;
  - un ✅ sur un texte à placeholders ([VAINQUEUR]…) ou sans média est bloqué et signalé,
    pas publié — publier un trou de texte fait plus de dégâts qu'un retard ;
  - idempotence par reviews/etat/<client>.json : une décision déjà appliquée à l'identique
    n'est pas rejouée, une décision changée est rejouée ;
  - un reel Instagram part en mode notification, jamais en publication automatique : la
    musique se pose dans Instagram, jamais dans le fichier (Maxime, 21/10/2026) ;
  - une image de feed hors de la plage 4:5 à 1.91:1 n'est pas publiée en automatique.

Secrets attendus : METRICOOL_USER_ID, METRICOOL_TOKEN. Dépendances : openpyxl, et Pillow
pour mesurer le format des images (sans Pillow, ce contrôle est sauté).
"""

import argparse
import hashlib
import io
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
from datetime import date, datetime, time as dtime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from publier import metricool  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
REVIEWS = ROOT / "reviews"
A_ENVOYER = REVIEWS / "a-envoyer"
RETOURS = REVIEWS / "retours"
ETAT = REVIEWS / "etat"

# Les trois décisions offertes au client. Le vide est la quatrième : « pas encore décidé »,
# et un post sans décision ne part jamais.
DECISION_OK = "✅ OK pour publier"
DECISION_CORRIGER = "✏️ À corriger (voir commentaire)"
DECISION_NON = "⛔ Ne pas publier"

JOURS = ["Lun", "Mar", "Mer", "Jeu", "Ven", "Sam", "Dim"]
RESEAUX = {"instagram": "IG", "facebook": "FB", "tiktok": "TikTok", "linkedin": "LinkedIn",
           "youtube": "YT", "twitter": "X", "pinterest": "Pinterest", "threads": "Threads",
           "gmb": "GMB"}
PLACEHOLDER = re.compile(r"\[[A-ZÀ-Ü0-9][A-ZÀ-Ü0-9 +/'-]*\]")
# Règle de rédaction de Maxime (mémoire redaction-sans-tirets) : aucun tiret cadratin ni
# double tiret dans un texte publié. Une légende qui en contient n'est pas publiée seule.
TIRET = re.compile(r"[—–]|--")
# Note de production laissée dans une légende (« 🎬 À tourner : … ») : le visuel attaché
# n'est qu'un provisoire, la vidéo reste à tourner. Jamais publié en automatique.
NOTE_INTERNE = re.compile(r"à tourner\s*:|a tourner\s*:|\[à faire\]|todo", re.IGNORECASE)
PARIS = ZoneInfo("Europe/Paris")
# Validation automatique : jusqu'où on regarde devant, et la marge avant publication
# (un post à moins de 30 minutes n'est plus touché, Metricool pourrait le publier deux fois).
HORIZON_AUTO_JOURS = 14
MARGE_AUTO = timedelta(minutes=30)
# Plage de format (largeur / hauteur) qu'accepte l'API Instagram pour une image de feed.
# Un 2:3 (0,67) passe dans l'appli mais est refusé ou recadré par l'API (audit du 21/10).
RATIO_FEED = (4 / 5, 1.91)
# Champs de couverture d'une vidéo, à recopier tels quels à chaque réécriture du post.
CHAMPS_COUVERTURE = ("videoCoverMilliseconds", "videoThumbnailUrl")


# ---------------------------------------------------------------- lecture Metricool

def semaine_lundi(texte=None):
    """Le lundi de la semaine visée. Sans argument : la semaine en cours."""
    if texte:
        d = date.fromisoformat(texte)
        if d.weekday() != 0:
            d -= timedelta(days=d.weekday())
            print(f"(la date donnée n'est pas un lundi — semaine du {d} retenue)")
        return d
    aujourd_hui = date.today()
    return aujourd_hui - timedelta(days=aujourd_hui.weekday())


def posts_semaine(blog_id, lundi):
    dimanche = lundi + timedelta(days=6)
    posts = metricool("/scheduler/posts", blogId=blog_id,
                      start=f"{lundi}T00:00:00", end=f"{dimanche}T23:59:59").get("data", [])
    return sorted(posts, key=lambda p: p.get("publicationDate", {}).get("dateTime", ""))


def type_du(post):
    """Le format lisible par le client. La donnée vit dans instagramData/facebookData/…

    Metricool renvoie un bloc *Data pour CHAQUE réseau, même ceux où le post ne part pas :
    une story Instagram arrive avec un twitterData de type POST avant son instagramData.
    On lit donc d'abord les blocs des réseaux réellement visés (providers), dans leur
    ordre ; les autres ne servent que de repli. Constaté le 05/10/2026 : toutes les
    stories de Pierre Huiban s'affichaient « Post » dans le tableau.
    """
    if len(post.get("media") or []) > 1:
        return f"Carrousel ({len(post['media'])} visuels)"
    vises = [f"{pr.get('network')}Data" for pr in post.get("providers") or [] if pr.get("network")]
    autres = [cle for cle in post if cle.endswith("Data") and cle not in vises]
    for cle in vises + autres:
        val = post.get(cle)
        if isinstance(val, dict) and val.get("type"):
            return {"POST": "Post", "REEL": "Reel", "STORY": "Story"}.get(val["type"], val["type"].title())
    return "Post"


def reseaux_du(post):
    abbr = [RESEAUX.get(pr.get("network"), pr.get("network", "?")) for pr in post.get("providers", [])]
    return " + ".join(abbr) or "?"


# ---------------------------------------------------------------- écriture du fichier

ENTETES = ["Réf", "Jour", "Heure", "Format", "Réseaux", "Visuel",
           "Légende", "Votre décision", "Votre commentaire"]
# Colonnes que le client édite. Tout le reste est verrouillé — le verrou tient dans Excel,
# Google Sheets l'ignore : c'est un garde-fou contre la fausse manœuvre, pas une serrure.
EDITABLES = {"C", "G", "H", "I"}
LARGEURS = {"A": 9, "B": 11, "C": 9, "D": 12, "E": 13, "F": 15, "G": 62, "H": 27, "I": 34}

NOIR = "FF111827"
GRIS_FOND = "FFF5F6F8"
JAUNE_EDIT = "FFFFF7DB"
VERT_OK = "FFE3F4E8"
ORANGE_CORRIGER = "FFFDEBD3"
ROUGE_NON = "FFFADBD8"
BORD = Border(*[Side(style="thin", color="FFD9DDE3")] * 4)


def _mode_demploi(wb, compte, lundi):
    ws = wb.active
    ws.title = "Mode d'emploi"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 100
    lignes = [
        (f"Vos publications : semaine du {lundi:%d/%m} au {lundi + timedelta(days=6):%d/%m}", "titre"),
        (compte["nom"], "sous-titre"),
        ("", None),
        ("Trois minutes suffisent :", "gras"),
        ("1.  Ouvrez l'onglet « Semaine du … » : chaque ligne est une publication prévue.", None),
        ("2.  Les colonnes JAUNES sont à vous : corrigez la légende directement dans la case,", None),
        ("     ajustez l'heure si besoin, et laissez-nous un commentaire si utile.", None),
        ("3.  Choisissez votre décision dans le menu de la colonne « Votre décision » :", None),
        (f"      {DECISION_OK} : la publication part automatiquement, plus rien à faire.", None),
        (f"      {DECISION_CORRIGER} : on la retravaille avec votre commentaire et on vous la représente.", None),
        (f"      {DECISION_NON} : elle est retirée de la semaine.", None),
        ("", None),
        ("Vos décisions sont relevées plusieurs fois par jour et prises en compte sous quelques", None),
        ("heures. Une ligne laissée sans décision reste en attente : elle ne sera pas publiée.", None),
        ("Les lignes « ⚠️ visuel à venir » reçoivent leur photo/vidéo avant publication.", None),
        ("", None),
        ("Merci de ne pas toucher aux colonnes grises ni à la colonne Réf : elles servent au suivi.", "petit"),
    ]
    for i, (texte, style) in enumerate(lignes, start=2):
        cellule = ws.cell(row=i, column=2, value=texte)
        if style == "titre":
            cellule.font = Font(bold=True, size=16, color=NOIR)
        elif style == "sous-titre":
            cellule.font = Font(bold=True, size=12, color="FF6B7280")
        elif style == "gras":
            cellule.font = Font(bold=True)
        elif style == "petit":
            cellule.font = Font(size=9, color="FF6B7280", italic=True)
    ws.protection.sheet = True
    return ws


def _hauteur(texte):
    """Une hauteur de ligne qui laisse la légende lisible sans dérouler."""
    lignes = sum(max(1, len(l) // 58 + 1) for l in (texte or "").splitlines()) or 1
    return min(15 * lignes + 6, 210)


def ecrire_fichier(compte, lundi, brouillons, chemin):
    wb = openpyxl.Workbook()
    _mode_demploi(wb, compte, lundi)

    ws = wb.create_sheet(f"Semaine du {lundi:%d-%m}")
    ws.freeze_panes = "A2"
    for col, largeur in LARGEURS.items():
        ws.column_dimensions[col].width = largeur

    for c, entete in enumerate(ENTETES, start=1):
        cellule = ws.cell(row=1, column=c, value=entete)
        cellule.font = Font(bold=True, color="FFFFFFFF")
        cellule.fill = PatternFill("solid", fgColor=NOIR)
        cellule.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22

    tech = wb.create_sheet("_tech")
    tech["A1"], tech["B1"] = "client", compte["slug"]
    tech["A2"], tech["B2"] = "semaine", lundi.isoformat()
    tech["A3"], tech["B3"] = "genere", datetime.now().isoformat(timespec="seconds")
    for c, entete in enumerate(["ref", "post_id", "blog_id", "date_time", "texte_original"], start=1):
        tech.cell(row=5, column=c, value=entete)

    for i, post in enumerate(brouillons, start=1):
        r = i + 1
        dt = datetime.fromisoformat(post["publicationDate"]["dateTime"])
        ref = f"{compte['prefixe']}-{i:02d}"
        media = post.get("media") or []
        valeurs = [ref, f"{JOURS[dt.weekday()]} {dt:%d/%m}", f"{dt:%H:%M}", type_du(post),
                   reseaux_du(post), "", post.get("text", ""), "", ""]
        for c, valeur in enumerate(valeurs, start=1):
            cellule = ws.cell(row=r, column=c, value=valeur)
            colonne = cellule.column_letter
            cellule.border = BORD
            cellule.alignment = Alignment(vertical="top",
                                          wrap_text=colonne in ("G", "I"),
                                          horizontal="center" if colonne in "ABCDEF" else "left")
            if colonne in EDITABLES:
                cellule.fill = PatternFill("solid", fgColor=JAUNE_EDIT)
                cellule.protection = Protection(locked=False)
            else:
                cellule.fill = PatternFill("solid", fgColor=GRIS_FOND)
        visuel = ws.cell(row=r, column=6)
        if media:
            visuel.value = "voir le visuel"
            visuel.hyperlink = media[0]
            visuel.font = Font(color="FF1155CC", underline="single")
        else:
            visuel.value = "⚠️ visuel à venir"
        ws.row_dimensions[r].height = _hauteur(post.get("text", ""))

        for c, valeur in enumerate([ref, post.get("id"), compte["blog_id"],
                                    post["publicationDate"]["dateTime"], post.get("text", "")], start=1):
            tech.cell(row=5 + i, column=c, value=valeur)

    fin = len(brouillons) + 1
    decisions = DataValidation(
        type="list", formula1=f'"{DECISION_OK},{DECISION_CORRIGER},{DECISION_NON}"',
        allow_blank=True, showErrorMessage=True, errorTitle="Choisir dans la liste",
        error="Choisissez une des trois décisions du menu, ou laissez la case vide.")
    ws.add_data_validation(decisions)
    decisions.add(f"H2:H{fin}")
    for fond, decision in ((VERT_OK, DECISION_OK), (ORANGE_CORRIGER, DECISION_CORRIGER),
                           (ROUGE_NON, DECISION_NON)):
        ws.conditional_formatting.add(
            f"A2:I{fin}",
            FormulaRule(formula=[f'$H2="{decision}"'],
                        fill=PatternFill("solid", bgColor=fond), stopIfTrue=True))

    ws.protection.sheet = True
    tech.sheet_state = "hidden"
    tech.protection.sheet = True

    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)


def semaines_a_venir(compte, aujourd_hui=None):
    """Les lundis à générer pour un compte lors de la génération planifiée du vendredi.

    `semaines_avance` (1 par défaut) : combien de semaines le client a ouvertes devant lui.
    À 2, le vendredi produit le lundi suivant ET celui d'après — Maxime le 05/10/2026 pour
    Pierre Huiban : « on aura deux semaines d'avance », valider le vendredi pour le lundi
    est trop juste. Le fichier du lundi suivant existe déjà en général (généré la semaine
    d'avant) : la routine du vendredi garde alors celui du Drive, où le client a peut-être
    déjà coché, et ne dépose que le nouveau.
    """
    aujourd_hui = aujourd_hui or date.today()
    prochain = aujourd_hui + timedelta(days=7 - aujourd_hui.weekday())
    return [prochain + timedelta(weeks=i) for i in range(max(1, int(compte.get("semaines_avance") or 1)))]


def fichiers_de(compte):
    """Les fichiers Drive que la relève lit pour un compte, sans doublon.

    `drive_fichiers` (liste) quand le client a plusieurs semaines ouvertes à la fois,
    `drive_fichier` (un seul) sinon ; les deux peuvent cohabiter le temps d'une bascule.
    """
    ids = list(compte.get("drive_fichiers") or []) + [compte.get("drive_fichier")]
    return [i for n, i in enumerate(ids) if i and i not in ids[:n]]


def generer_compte(compte, lundi, dry_run=False):
    """Écrit le fichier de review d'un compte. Retourne le nombre de brouillons couverts."""
    posts = posts_semaine(compte["blog_id"], lundi)
    brouillons = [p for p in posts if p.get("draft")]
    if not brouillons:
        print(f"= {compte['slug']} : aucun brouillon du {lundi} au {lundi + timedelta(days=6)} — pas de fichier")
        return 0

    a_completer = [p for p in brouillons if PLACEHOLDER.search(p.get("text", ""))]
    sans_media = [p for p in brouillons if not p.get("media")]
    if dry_run:
        print(f"[dry-run] {compte['slug']} : {len(brouillons)} brouillon(s), "
              f"{len(a_completer)} placeholder(s), {len(sans_media)} sans média")
        for p in brouillons:
            print(f"           └ {p['publicationDate']['dateTime']} {type_du(p)} : {p.get('text', '')[:60]}…")
        return len(brouillons)

    chemin = A_ENVOYER / f"{compte['slug']}-{lundi}.xlsx"
    ecrire_fichier(compte, lundi, brouillons, chemin)
    print(f"✓ {compte['slug']} : {len(brouillons)} brouillon(s) → {chemin.relative_to(ROOT)}")
    if a_completer:
        print(f"  ⚠ {len(a_completer)} texte(s) à placeholders — à compléter AVANT que le client valide")
    if sans_media:
        print(f"  ⚠ {len(sans_media)} post(s) sans média — un ✅ restera bloqué tant que le visuel manque")
    return len(brouillons)


# ---------------------------------------------------------------- lecture des retours

def telecharger_review(drive_id):
    """L'export public du fichier de review, sans identifiant Google.

    Deux formes possibles selon ce que le connecteur a créé au dépôt : un Google Sheet
    (conversion par défaut) s'exporte en xlsx via docs.google.com, un .xlsx resté brut se
    télécharge par le lien direct. On tente dans cet ordre. Un xlsx est un zip : la
    signature PK distingue un vrai fichier d'une page HTML « demande d'accès ».
    """
    urls = [f"https://docs.google.com/spreadsheets/d/{drive_id}/export?format=xlsx",
            f"https://drive.google.com/uc?export=download&id={drive_id}"]
    for url in urls:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                octets = r.read()
        except urllib.error.HTTPError:
            continue
        if octets[:2] == b"PK":
            return octets
    raise RuntimeError("export impossible — le fichier n'est probablement plus partagé "
                       "« tous les utilisateurs disposant du lien »")


def _iso(valeur):
    """Une date de _tech telle qu'écrite. La conversion en Google Sheet retype parfois nos
    chaînes ISO en vraies dates : au retour, on re-normalise plutôt que de planter."""
    if isinstance(valeur, datetime):
        return valeur.strftime("%Y-%m-%dT%H:%M:%S")
    return str(valeur)


def lire_review(source):
    """Extrait (meta, lignes) d'un fichier de review, quel que soit son canal de retour."""
    contenu = io.BytesIO(source) if isinstance(source, bytes) else source
    wb = openpyxl.load_workbook(contenu, data_only=True)
    if "_tech" not in wb.sheetnames:
        raise RuntimeError("onglet _tech absent — ce fichier n'a pas été généré par review.py")
    tech = wb["_tech"]
    meta = {"client": tech["B1"].value, "semaine": _iso(tech["B2"].value)[:10]}
    originaux = {}
    for ligne in tech.iter_rows(min_row=6, max_col=5, values_only=True):
        if ligne[0]:
            originaux[str(ligne[0])] = {"post_id": ligne[1], "blog_id": ligne[2],
                                        "date_time": _iso(ligne[3]), "texte": ligne[4] or ""}
    feuille = next((ws for ws in wb.worksheets if ws.title.startswith("Semaine")), None)
    if feuille is None:
        raise RuntimeError("onglet « Semaine du … » introuvable")
    lignes = []
    for valeurs in feuille.iter_rows(min_row=2, max_col=9, values_only=True):
        ref = str(valeurs[0]).strip() if valeurs[0] else None
        if not ref:
            continue
        lignes.append({"ref": ref, "heure": valeurs[2], "legende": valeurs[6],
                       "decision": valeurs[7], "commentaire": valeurs[8],
                       "original": originaux.get(ref)})
    return meta, lignes


def normaliser_texte(texte):
    return "\n".join(l.rstrip() for l in str(texte or "").replace("\r\n", "\n").split("\n")).strip()


def normaliser_heure(cellule):
    """« 18:00 », 18h, 18 h 30, « 18 » nu, une heure Excel ou un sériel Sheets → « HH:MM ».

    Une heure illisible ou impossible (25h…) vaut None = « heure inchangée » : on garde le
    créneau d'origine plutôt que d'inventer, l'écart se voit dans le fichier au pire.
    """
    if cellule is None or str(cellule).strip() == "":
        return None
    if isinstance(cellule, datetime):
        return cellule.strftime("%H:%M")
    if isinstance(cellule, dtime):
        return cellule.strftime("%H:%M")
    if isinstance(cellule, (int, float)) and not isinstance(cellule, bool):
        valeur = float(cellule)
        if 0 <= valeur < 1:  # sériel : fraction de jour
            minutes = round(valeur * 24 * 60)
            return f"{minutes // 60:02d}:{minutes % 60:02d}"
        if valeur.is_integer() and 0 <= valeur < 24:  # « 18 » tapé nu, devenu nombre
            return f"{int(valeur):02d}:00"
        return None
    texte = str(cellule)
    m = re.search(r"(\d{1,2})\s*[hH:.]\s*(\d{1,2})?", texte)
    if not m:
        m = re.fullmatch(r"\s*(\d{1,2})\s*", texte)
        if m and int(m.group(1)) < 24:
            return f"{int(m.group(1)):02d}:00"
        return None
    heures, minutes = int(m.group(1)), int(m.group(2) or 0)
    if heures > 23 or minutes > 59:
        return None
    return f"{heures:02d}:{minutes:02d}"


def lire_decision(cellule):
    """Décision d'une ligne, tolérante aux clients qui tapent au lieu de choisir dans le menu.

    L'ordre des tests compte : « ne pas publier » contient « publier », le refus se teste
    donc avant l'accord. Tout ce qui n'est pas reconnu devient INCONNU : signalé à Maxime,
    jamais interprété — on ne publie pas sur une devinette.
    """
    if cellule is None or str(cellule).strip() == "":
        return None
    brut = str(cellule)
    plat = unicodedata.normalize("NFKD", brut.strip().lower())
    plat = "".join(c for c in plat if not unicodedata.combining(c))
    if "⛔" in brut or "pas publier" in plat or plat in {"non", "refuse", "no", "stop"}:
        return "NON"
    if "✅" in brut or "ok" in plat or "valid" in plat or "publier" in plat or plat in {"oui", "go", "yes"}:
        return "OK"
    if "✏" in brut or "corrig" in plat or "modif" in plat or "revoir" in plat or "retravail" in plat:
        return "CORRIGER"
    return "INCONNU"


# ---------------------------------------------------------------- application Metricool

def vise_instagram(post):
    return any(pr.get("network") == "instagram" for pr in post.get("providers") or [])


def est_reel_instagram_sans_son(post):
    """Un reel Instagram dont la musique reste à poser dans Instagram.

    Règle de Maxime (21/10/2026, déjà dite avant) : la musique d'un reel se prend dans la
    bibliothèque d'Instagram et ne s'intègre jamais au fichier, pour que le reel porte un
    son Instagram et non un « audio d'origine ». Seule exception : un reel dont Metricool
    pose lui-même le son Instagram (instagramData.audioConfiguration).
    """
    donnees = post.get("instagramData") or {}
    return (vise_instagram(post) and donnees.get("type") in ("REEL", "TRIAL_REEL")
            and not donnees.get("audioConfiguration"))


def corps_de(post, texte, date_time, publier_auto, notification=False, contenu_ia=False):
    """Le corps d'écriture d'un post, à partir de sa version lue dans Metricool.

    Même socle que publier.py (text, publicationDate, draft, autoPublish, media, providers)
    plus les blocs *Data : c'est là que vit le type POST/REEL/STORY, le perdre transforme
    un reel en post feed. La couverture d'une vidéo se recopie aussi : la perdre remet la
    première image en couverture.

    `notification` : le post sort du brouillon sans publication automatique. Metricool
    envoie alors une notification à Maxime à l'heure prévue, il ajoute la musique dans
    Instagram et publie. `contenu_ia` : compte dont les visuels sont générés par IA
    (Run On Invest) ; la mention « contenu IA » d'Instagram est alors toujours cochée.
    """
    corps = {
        "text": texte,
        "publicationDate": {"dateTime": date_time,
                            "timezone": post.get("publicationDate", {}).get("timezone", "Europe/Paris")},
        "draft": not publier_auto,
        "autoPublish": publier_auto and not notification,
        "media": post.get("media") or [],
        "providers": [{"network": pr["network"]} for pr in post.get("providers", []) if pr.get("network")],
    }
    for cle, valeur in post.items():
        if cle.endswith("Data") and isinstance(valeur, dict):
            corps[cle] = valeur
    for cle in CHAMPS_COUVERTURE:
        if post.get(cle) is not None:
            corps[cle] = post[cle]
    if contenu_ia and vise_instagram(post):
        corps["instagramData"] = {**corps.get("instagramData", {}), "isAiGenerated": True}
    if post.get("firstCommentText"):
        corps["firstCommentText"] = post["firstCommentText"]
    return corps


def dimensions_image(url, plafond=2_000_000):
    """(largeur, hauteur) d'une image distante, en ne lisant que son en-tête ; None si
    l'image est illisible, inaccessible, ou si Pillow manque."""
    try:
        from PIL import ImageFile
    except ImportError:
        return None
    analyseur = ImageFile.Parser()
    try:
        with urllib.request.urlopen(url, timeout=20) as reponse:
            lu = 0
            while lu < plafond:
                bloc = reponse.read(65536)
                if not bloc:
                    break
                lu += len(bloc)
                analyseur.feed(bloc)
                if analyseur.image is not None:
                    return analyseur.image.size
    except Exception:
        return None
    return analyseur.image.size if analyseur.image is not None else None


def format_hors_feed(post, mesurer):
    """Le format (largeur / hauteur) de la première image hors plage feed, sinon None.

    Ne regarde que les posts feed et carrousels visant Instagram ; une story ou un reel
    a sa propre plage. Une image qu'on n'arrive pas à mesurer ne bloque pas.
    """
    if not vise_instagram(post) or type_du(post) in ("Story", "Reel"):
        return None
    for url in post.get("media") or []:
        if re.search(r"\.(mp4|mov|m4v)(\?|$)", url, re.IGNORECASE):
            continue
        taille = mesurer(url)
        if taille and taille[1]:
            ratio = taille[0] / taille[1]
            if not RATIO_FEED[0] - 0.01 <= ratio <= RATIO_FEED[1] + 0.01:
                return ratio
    return None


def passer_en_publication(blog_id, post, texte, date_time, contenu_ia=False):
    """Applique texte/horaire et sort le post du brouillon. Retourne (résultat, post_id).

    Un reel Instagram sans son Instagram posé par Metricool sort en mode notification
    (draft false, autoPublish false) : Maxime ajoute la musique dans Instagram et publie.

    ⚠️ Constaté au premier test réel (post 355355797 → 359735766, OTEA, 08/08/2026) : le
    PUT de l'API ne modifie PAS l'entrée existante — il crée une NOUVELLE entrée de
    planning (nouvel id, même uuid) avec les nouvelles valeurs, et laisse l'ancien
    brouillon en place. D'où la marche en trois temps, chacun vérifié sur l'état réel :

      1. la version publiable existe-t-elle déjà ? (un run précédent a pu la poser sans
         finir son ménage) — si oui, ne PAS ré-écrire : un second PUT publierait en double ;
      2. sinon PUT, puis relire le jour pour trouver la version réellement en place ;
      3. ménage : retirer l'ancien id s'il est resté en brouillon, et les éventuels
         doublons publiables identiques. Seul cas de suppression automatique du système,
         strictement borné : l'entrée périmée d'un post que le client vient de valider —
         sans ce ménage, la review suivante re-proposerait un post déjà validé.

    Si l'API refuse le PUT : suppression puis recréation. L'ordre (inverse de publier.py)
    est voulu : ici le post recréé part en publication automatique, un doublon publierait
    DEUX fois chez le client. Mieux vaut un trou signalé qu'un doublon publié.
    """
    ancien_id = post["id"]
    jour = date_time[:10]
    notification = est_reel_instagram_sans_son(post)
    libelle = ("notification programmée, musique à poser dans Instagram" if notification
               else "publication programmée")
    corps = corps_de(post, texte, date_time, True, notification=notification,
                     contenu_ia=contenu_ia)

    def etat_du_jour():
        posts_jour = metricool("/scheduler/posts", blogId=blog_id,
                               start=f"{jour}T00:00:00", end=f"{jour}T23:59:59").get("data", [])
        publiables = [p for p in posts_jour
                      if not p.get("draft") and bool(p.get("autoPublish")) != notification
                      and p.get("publicationDate", {}).get("dateTime") == date_time
                      and normaliser_texte(p.get("text")) == normaliser_texte(texte)]
        reste = any(p.get("id") == ancien_id and p.get("draft") for p in posts_jour)
        return publiables, reste

    publiables, reste = etat_du_jour()
    if not publiables:
        try:
            metricool(f"/scheduler/posts/{ancien_id}", "PUT", corps, blogId=blog_id)
        except RuntimeError as exc:
            print(f"    (mise à jour refusée — {exc} — suppression puis recréation)")
            metricool(f"/scheduler/posts/{ancien_id}", "DELETE", blogId=blog_id)
            cree = metricool("/scheduler/posts", "POST", corps, blogId=blog_id)
            nouveau = cree.get("data", {}).get("id")
            return f"{libelle} (recréé, nouvel id {nouveau})", nouveau
        publiables, reste = etat_du_jour()
    if not publiables:
        raise RuntimeError(f"aucune version publiable en place après la mise à jour du "
                           f"post {ancien_id} — vérifier dans Metricool")

    version = publiables[0]
    for doublon in publiables[1:]:
        metricool(f"/scheduler/posts/{doublon['id']}", "DELETE", blogId=blog_id)
    if reste and version["id"] != ancien_id:
        metricool(f"/scheduler/posts/{ancien_id}", "DELETE", blogId=blog_id)
        return (f"{libelle} (version {version['id']}, ancien brouillon "
                f"{ancien_id} retiré)"), version["id"]
    return f"{libelle} (version {version['id']})", version["id"]


def retrouver_post(posts, original):
    """Le brouillon visé, par id d'abord, par créneau horaire sinon.

    Le repêchage par horaire n'est pas un luxe : les ids changent quand un brouillon est
    recréé (correction d'horaire, --remplacer de publier.py). Même clé que
    carrousels_existants : le créneau, pas le texte.
    """
    for p in posts:
        if p.get("id") == original["post_id"]:
            return p
    candidats = [p for p in posts
                 if p.get("publicationDate", {}).get("dateTime") == original["date_time"]
                 and p.get("draft")]
    return candidats[0] if len(candidats) == 1 else None


def charger_etat(slug):
    chemin = ETAT / f"{slug}.json"
    if chemin.exists():
        return json.loads(chemin.read_text(encoding="utf-8"))
    return {"_doc": "Journal des décisions client appliquées par scripts/review.py — écrit par "
                    "le workflow « Review hebdo client », ne pas éditer à la main. Sert "
                    "d'idempotence (une décision déjà appliquée à l'identique n'est pas rejouée) "
                    "et de matière au récap du lundi.",
            "semaines": {}}


def empreinte(decision, legende, heure):
    brut = f"{decision}|{normaliser_texte(legende)}|{heure or ''}"
    return hashlib.sha256(brut.encode()).hexdigest()[:16]


def appliquer_compte(compte, dry_run=False):
    """Traite toutes les sources de retour d'un compte. Retourne les compteurs."""
    slug = compte["slug"]
    sources, illisibles = [], 0
    for drive_id in fichiers_de(compte):
        try:
            sources.append((f"Drive {drive_id[:8]}…", telecharger_review(drive_id)))
        except Exception as exc:
            # Un fichier illisible n'empêche pas de relever l'autre semaine ouverte.
            print(f"✗ {slug} : export Drive impossible ({drive_id}) — {exc}", file=sys.stderr)
            annoter("warning", f"{slug} : fichier de review illisible ({exc}). Il n'est pas "
                               f"relevé tant que le dossier Drive n'est pas partagé par lien.")
            illisibles += 1
    for chemin in sorted(RETOURS.glob(f"{slug}*.xlsx")):
        sources.append((chemin.name, chemin))
    if not sources:
        if illisibles:
            return {"erreurs": illisibles}
        print(f"= {slug} : aucun fichier de review à lire (drive_fichier vide, aucun retour déposé)")
        return {}

    etat = charger_etat(slug)
    compteurs = {"publies": 0, "a_corriger": 0, "refuses": 0, "erreurs": illisibles,
                 "sans_decision": 0, "en_attente": 0}
    posts_cache = {}  # par semaine : deux fichiers ouverts = deux semaines de posts
    touches = False

    for nom_source, source in sources:
        try:
            meta, lignes = lire_review(source)
        except Exception as exc:
            print(f"✗ {slug} : {nom_source} illisible — {exc}", file=sys.stderr)
            compteurs["erreurs"] += 1
            continue
        if meta["client"] != slug:
            print(f"! {slug} : {nom_source} appartient à « {meta['client']} », ignoré")
            continue
        semaine = str(meta["semaine"])
        journal = etat["semaines"].setdefault(semaine, {})
        print(f"{compte['nom']} — semaine du {semaine} (source : {nom_source})")

        for ligne in lignes:
            ref, original = ligne["ref"], ligne["original"]
            decision = lire_decision(ligne["decision"])
            legende = normaliser_texte(ligne["legende"])
            heure = normaliser_heure(ligne["heure"])
            commentaire = str(ligne["commentaire"]).strip() if ligne["commentaire"] else None

            if original is None:
                print(f"  ! {ref} : référence inconnue de l'onglet _tech — ligne ignorée")
                compteurs["erreurs"] += 1
                continue

            legende_modifiee = legende != normaliser_texte(original["texte"])
            heure_originale = original["date_time"][11:16]
            heure_modifiee = heure is not None and heure != heure_originale

            if decision is None:
                if legende_modifiee or commentaire:
                    compteurs["sans_decision"] += 1
                    print(f"  ℹ {ref} : ligne retouchée mais AUCUNE décision — rien fait, à trancher")
                continue

            cle = empreinte(decision, legende, heure)
            deja = journal.get(ref, {})
            if deja.get("hash") == cle and not deja.get("en_attente"):
                continue  # déjà appliqué à l'identique — sauf ✅ retenu, qui se retente

            entree = {"decision": decision, "hash": cle,
                      "quand": datetime.now().isoformat(timespec="seconds"),
                      "commentaire": commentaire, "legende_modifiee": legende_modifiee,
                      "heure": heure if heure_modifiee else None,
                      "post_id": original["post_id"], "resultat": None}

            if decision == "INCONNU":
                entree["resultat"] = f"décision illisible : « {ligne['decision']} » — à clarifier"
                print(f"  ? {ref} : {entree['resultat']}")
            elif decision == "NON":
                entree["resultat"] = "refusé — laissé en brouillon, à retirer à la main"
                compteurs["refuses"] += 1
                print(f"  ⛔ {ref} : refusé" + (f" — « {commentaire} »" if commentaire else ""))
            elif decision == "CORRIGER":
                entree["resultat"] = "à corriger puis re-proposer"
                compteurs["a_corriger"] += 1
                print(f"  ✏️ {ref} : à corriger — « {commentaire or 'sans commentaire'} »")
            else:  # OK
                if semaine not in posts_cache:
                    lundi = date.fromisoformat(semaine)
                    posts_cache[semaine] = posts_semaine(compte["blog_id"], lundi)
                post = retrouver_post(posts_cache[semaine], original)
                nouvelle_date = f"{original['date_time'][:11]}{heure}:00" if heure_modifiee \
                    else original["date_time"]
                # Une case légende vidée est presque toujours une fausse manœuvre : on garde
                # le texte d'origine plutôt que de publier une légende vide.
                texte_final = legende if (legende_modifiee and legende) else original["texte"]

                if post is None:
                    entree["resultat"] = "post introuvable dans Metricool (supprimé ou déplacé hors semaine)"
                    compteurs["erreurs"] += 1
                    print(f"  ✗ {ref} : {entree['resultat']}")
                elif not post.get("draft"):
                    entree["resultat"] = "déjà validé directement dans Metricool — rien à faire"
                    print(f"  = {ref} : {entree['resultat']}")
                elif PLACEHOLDER.search(texte_final):
                    # ✅ retenu, pas refusé : l'accord du client reste acquis, la relève
                    # retente à chaque passage et publie dès que le texte est complété.
                    entree["resultat"] = ("✅ retenu — texte à placeholders, publication en attente : "
                                          "compléter le texte et la relève suivante publiera")
                    entree["en_attente"] = True
                    compteurs["en_attente"] += 1
                    print(f"  ⏳ {ref} : {entree['resultat']}")
                elif not post.get("media"):
                    entree["resultat"] = ("✅ retenu — aucun média attaché, publication en attente : "
                                          "ajouter le visuel dans Metricool et la relève suivante publiera")
                    entree["en_attente"] = True
                    compteurs["en_attente"] += 1
                    print(f"  ⏳ {ref} : {entree['resultat']}")
                elif dry_run:
                    detail = [d for d, actif in (("légende", legende_modifiee),
                                                 (f"heure → {heure}", heure_modifiee)) if actif]
                    print(f"  [dry-run] {ref} : passerait en publication auto"
                          + (f" ({', '.join(detail)})" if detail else ""))
                    continue
                else:
                    try:
                        resultat, post_id = passer_en_publication(
                            compte["blog_id"], post, texte_final, nouvelle_date,
                            contenu_ia=compte.get("contenu_ia", False))
                        entree["resultat"], entree["post_id"] = resultat, post_id
                        compteurs["publies"] += 1
                        detail = " — légende du client" if legende_modifiee else ""
                        print(f"  ✓ {ref} : {resultat}{detail}")
                        posts_cache.pop(semaine, None)  # cache invalide après écriture
                    except Exception as exc:
                        entree["resultat"] = f"échec Metricool : {exc}"
                        compteurs["erreurs"] += 1
                        print(f"  ✗ {ref} : {exc}", file=sys.stderr)

            if not dry_run:
                journal[ref] = entree
                touches = True

    if touches:
        ETAT.mkdir(parents=True, exist_ok=True)
        (ETAT / f"{slug}.json").write_text(
            json.dumps(etat, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return compteurs


# ---------------------------------------------------------------- orchestration

def annoter(niveau, message):
    """Annotation GitHub Actions (warning/error) : lisible dans la liste des runs sans ouvrir
    les logs. Trois relèves par jour en rouge sans dire pourquoi, personne ne les lit plus.
    Sans effet hors d'un runner."""
    if os.environ.get("GITHUB_ACTIONS"):
        print(f"::{niveau}::{message}")


def resumer(lignes):
    """Résumé Markdown du run, dans l'onglet Summary de GitHub Actions. Sans effet ailleurs."""
    chemin = os.environ.get("GITHUB_STEP_SUMMARY")
    if chemin:
        with open(chemin, "a", encoding="utf-8") as f:
            f.write("\n".join(lignes) + "\n")

def a_valider_auto(posts, maintenant, horizon_jours=HORIZON_AUTO_JOURS, medias_retenus=(),
                   mesurer=None):
    """Tri des brouillons d'un compte en validation automatique.

    Retourne (a_publier, retenus), retenus étant une liste de (post, raison). Seuls les
    brouillons à venir (au-delà de la marge, en deçà de l'horizon) sont regardés. Les
    garde-fous de la review restent : texte à placeholders ou sans média = retenu. S'ajoute
    le tiret cadratin, qui ne compte pas pour une story (Instagram n'affiche pas son texte),
    et les visuels listés dans `medias_retenus` du compte (morceaux d'URL : visuels IA
    interdits sur un compte qui a de vraies photos, règle du 21/08/2026). Avec `mesurer`
    (url → (largeur, hauteur)), une image de feed hors de la plage Instagram est retenue.
    """
    limite = maintenant + timedelta(days=horizon_jours)
    a_publier, retenus = [], []
    for post in posts:
        if not post.get("draft"):
            continue
        quand = datetime.fromisoformat(post["publicationDate"]["dateTime"])
        if quand <= maintenant + MARGE_AUTO or quand > limite:
            continue
        texte = post.get("text") or ""
        if PLACEHOLDER.search(texte):
            retenus.append((post, "texte à placeholders"))
        elif not post.get("media"):
            retenus.append((post, "aucun média attaché"))
        elif any(m in url for url in post["media"] for m in medias_retenus):
            retenus.append((post, "visuel IA à remplacer par une vraie photo"))
        elif NOTE_INTERNE.search(texte):
            retenus.append((post, "note de tournage dans la légende (vidéo à tourner)"))
        elif TIRET.search(texte) and type_du(post) != "Story":
            retenus.append((post, "tiret cadratin dans la légende (règle de rédaction)"))
        elif mesurer and (ratio := format_hors_feed(post, mesurer)) is not None:
            retenus.append((post, f"image hors du format feed d'Instagram ({ratio:.2f}, "
                                  f"à recadrer entre 4:5 et 1.91:1)"))
        else:
            a_publier.append(post)
    return a_publier, retenus


def valider_auto(compte, dry_run=False, maintenant=None):
    """Compte dont Maxime est lui-même le client : chaque brouillon vaut ✅.

    Décision de Maxime du 07/10/2026 pour OTEA Production et Otea Filmakers : « tu peux
    considérer que c'est validé à chaque fois, car c'est moi ». Étendue le 21/10/2026 à
    Pierre Huiban, qui ne validait pas ses tableaux, et à Run On Invest (contenu_ia : la
    mention « contenu IA » est cochée à chaque passage). Pas de tableau ; à chaque
    relève, les brouillons des deux semaines à venir qui passent les garde-fous partent en
    publication automatique, par le même chemin qu'un ✅ client (passer_en_publication).
    """
    slug = compte["slug"]
    maintenant = maintenant or datetime.now(PARIS).replace(tzinfo=None)
    debut, fin = maintenant.date(), maintenant.date() + timedelta(days=HORIZON_AUTO_JOURS)
    posts = metricool("/scheduler/posts", blogId=compte["blog_id"],
                      start=f"{debut}T00:00:00", end=f"{fin}T23:59:59").get("data", [])
    a_publier, retenus = a_valider_auto(posts, maintenant,
                                        medias_retenus=compte.get("medias_retenus") or (),
                                        mesurer=dimensions_image)
    compteurs = {"publies": 0, "a_corriger": 0, "refuses": 0, "erreurs": 0,
                 "sans_decision": 0, "en_attente": len(retenus)}
    print(f"{compte['nom']} — validation automatique : {len(a_publier)} à publier, "
          f"{len(retenus)} retenu(s)")
    etat = charger_etat(slug)
    journal = etat.setdefault("auto", {})
    for post, raison in retenus:
        print(f"  ⏳ {post['publicationDate']['dateTime']} {type_du(post)} : {raison}")
    for post in a_publier:
        quand = post["publicationDate"]["dateTime"]
        if dry_run:
            print(f"  [dry-run] {quand} {type_du(post)} : passerait en publication auto")
            continue
        try:
            resultat, post_id = passer_en_publication(compte["blog_id"], post,
                                                      post.get("text") or "", quand,
                                                      contenu_ia=compte.get("contenu_ia", False))
            compteurs["publies"] += 1
            journal[quand + "|" + str(post["id"])] = {
                "quand": datetime.now().isoformat(timespec="seconds"),
                "type": type_du(post), "resultat": resultat, "post_id": post_id}
            print(f"  ✓ {quand} {type_du(post)} : {resultat}")
        except Exception as exc:
            compteurs["erreurs"] += 1
            print(f"  ✗ {quand} : {exc}", file=sys.stderr)
    if compteurs["publies"] and not dry_run:
        ETAT.mkdir(parents=True, exist_ok=True)
        (ETAT / f"{slug}.json").write_text(
            json.dumps(etat, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return compteurs


def comptes_auto(filtre=None):
    config = json.loads((REVIEWS / "config.json").read_text(encoding="utf-8"))
    return [c for c in config["comptes"] if c.get("validation_auto")
            and (not filtre or c["slug"] == filtre)]


def charger_comptes(filtre=None):
    config = json.loads((REVIEWS / "config.json").read_text(encoding="utf-8"))
    comptes = [c for c in config["comptes"] if c.get("actif")]
    if filtre:
        comptes = [c for c in comptes if c["slug"] == filtre]
        if not comptes:
            raise SystemExit(f"client inconnu ou inactif : {filtre}")
    return comptes


def main():
    ap = argparse.ArgumentParser(description="Review hebdo client — fichier Excel ↔ Metricool.")
    sous = ap.add_subparsers(dest="commande", required=True)
    g = sous.add_parser("generer", help="écrit le fichier Excel de la semaine par client")
    g.add_argument("--semaine", help="lundi visé (YYYY-MM-DD). Vide = semaine en cours")
    g.add_argument("--a-venir", action="store_true",
                   help="génération planifiée du vendredi : le lundi suivant, et les semaines "
                        "d'après selon semaines_avance de chaque compte")
    a = sous.add_parser("appliquer", help="lit les retours clients et pousse dans Metricool")
    for p in (g, a):
        p.add_argument("--client", help="ne traiter qu'un compte, ex. le-loft")
        p.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    absents = [v for v in ("METRICOOL_USER_ID", "METRICOOL_TOKEN") if not os.environ.get(v)]
    if absents:
        raise SystemExit(f"secrets manquants : {', '.join(absents)}")

    if args.commande == "appliquer" and args.client and comptes_auto(args.client):
        comptes = []  # compte en validation automatique : pas de tableau à relever
    else:
        comptes = charger_comptes(args.client)
    if args.commande == "generer":
        if args.a_venir:
            plan = [(c, lundi) for c in comptes for lundi in semaines_a_venir(c)]
            print(f"Review des semaines à venir — {len(comptes)} compte(s), {len(plan)} fichier(s) visé(s)\n")
        else:
            lundi = semaine_lundi(args.semaine)
            plan = [(c, lundi) for c in comptes]
            print(f"Review de la semaine du {lundi} au {lundi + timedelta(days=6)} — {len(comptes)} compte(s)\n")
        total = sum(generer_compte(c, lundi, args.dry_run) for c, lundi in plan)
        verbe = "couverts (rien n'a été écrit)" if args.dry_run else "couverts par les fichiers générés"
        print(f"\n{total} brouillon(s) {verbe}.")
        return 0

    totaux, resume = {}, ["## Relève de la review client", "", "| Compte | Publiés | À corriger | Refusés | En attente | Sans décision | Erreurs |",
                          "|---|---:|---:|---:|---:|---:|---:|"]
    for compte, traiter in [(c, appliquer_compte) for c in comptes] + \
            [(c, valider_auto) for c in comptes_auto(args.client)]:
        compteurs = traiter(compte, args.dry_run)
        for cle, valeur in compteurs.items():
            totaux[cle] = totaux.get(cle, 0) + valeur
        resume.append(f"| {compte['nom']} | " + " | ".join(
            str(compteurs.get(k, 0)) for k in ("publies", "a_corriger", "refuses", "en_attente",
                                               "sans_decision", "erreurs")) + " |")
    bilan = (f"Bilan : {totaux.get('publies', 0)} publié(s), {totaux.get('a_corriger', 0)} à corriger, "
             f"{totaux.get('refuses', 0)} refusé(s), {totaux.get('en_attente', 0)} ✅ en attente "
             f"(média/placeholders), {totaux.get('sans_decision', 0)} sans décision, "
             f"{totaux.get('erreurs', 0)} erreur(s).")
    print("\n" + bilan + ("" if args.dry_run else " Détail durable : reviews/etat/."))
    resumer(resume + ["", bilan])
    return 1 if totaux.get("erreurs") else 0


if __name__ == "__main__":
    sys.exit(main())

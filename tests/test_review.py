"""Le moteur de la review hebdo, testé sans réseau.

Les fonctions pures d'abord (lecture des décisions, des heures, des placeholders), puis
l'aller-retour complet : un fichier écrit par `ecrire_fichier` doit être relu à
l'identique par `lire_review`. C'est ce contrat qui permet au client de corriger dans
Google Sheets sans que la relève perde le rapprochement avec Metricool.
"""

import sys
from datetime import date, datetime, time as dtime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import review  # noqa: E402


# ---------------------------------------------------------------- semaine

def test_semaine_lundi_ramene_au_lundi():
    assert review.semaine_lundi("2026-09-09") == date(2026, 9, 7)
    assert review.semaine_lundi("2026-09-07") == date(2026, 9, 7)
    assert review.semaine_lundi("2026-09-13") == date(2026, 9, 7)


def test_semaine_lundi_sans_argument_est_un_lundi():
    assert review.semaine_lundi().weekday() == 0


# ---------------------------------------------------------------- lecture des posts

def test_type_du_distingue_carrousel_reel_story_post():
    assert review.type_du({"media": ["a", "b", "c"]}) == "Carrousel (3 visuels)"
    assert review.type_du({"media": ["a"], "instagramData": {"type": "REEL"}}) == "Reel"
    assert review.type_du({"media": [], "facebookData": {"type": "STORY"}}) == "Story"
    assert review.type_du({"media": ["a"], "instagramData": {"type": "POST"}}) == "Post"
    assert review.type_du({}) == "Post"


def test_type_du_lit_le_reseau_vise_avant_les_autres_blocs():
    # Metricool renvoie un twitterData de type POST même pour une story Instagram seule.
    story = {"media": ["a"], "providers": [{"network": "instagram"}],
             "twitterData": {"type": "POST"}, "instagramData": {"type": "STORY"}}
    assert review.type_du(story) == "Story"
    reel = {"media": [], "providers": [{"network": "instagram"}, {"network": "facebook"}],
            "twitterData": {"type": "POST"}, "facebookData": {"type": "POST"},
            "instagramData": {"type": "REEL"}}
    assert review.type_du(reel) == "Reel"


def test_reseaux_du_abrege_et_tolere_l_inconnu():
    post = {"providers": [{"network": "instagram"}, {"network": "facebook"}]}
    assert review.reseaux_du(post) == "IG + FB"
    assert review.reseaux_du({"providers": [{"network": "bluesky"}]}) == "bluesky"
    assert review.reseaux_du({}) == "?"


@pytest.mark.parametrize("texte, attendu", [
    ("Bravo à [VAINQUEUR] pour sa victoire", True),
    ("Résultat : [À COMPLÉTER]", True),
    ("[STAT] % des golfeurs", True),
    ("un [lien] en minuscules n'est pas un placeholder", False),
    ("aucun crochet ici", False),
    ("", False),
])
def test_placeholder(texte, attendu):
    assert bool(review.PLACEHOLDER.search(texte)) is attendu


# ---------------------------------------------------------------- heures

@pytest.mark.parametrize("cellule, attendu", [
    ("18:00", "18:00"),
    ("18h", "18:00"),
    ("18 h 30", "18:30"),
    ("9H05", "09:05"),
    ("18", "18:00"),
    (18, "18:00"),
    (0.75, "18:00"),                      # sériel Google Sheets : fraction de jour
    (dtime(9, 5), "09:05"),
    (datetime(2026, 9, 7, 18, 30), "18:30"),
    ("25h", None),                        # impossible : on garde l'heure d'origine
    ("18:75", None),
    ("bientôt", None),
    (None, None),
    ("   ", None),
    (True, None),
    (42, None),
])
def test_normaliser_heure(cellule, attendu):
    assert review.normaliser_heure(cellule) == attendu


# ---------------------------------------------------------------- décisions

@pytest.mark.parametrize("cellule, attendu", [
    (review.DECISION_OK, "OK"),
    (review.DECISION_CORRIGER, "CORRIGER"),
    (review.DECISION_NON, "NON"),
    ("ok", "OK"),
    ("OUI", "OK"),
    ("Validé", "OK"),
    ("publier", "OK"),
    ("ne pas publier", "NON"),            # contient « publier » : le refus prime
    ("Non", "NON"),
    ("STOP", "NON"),
    ("à corriger", "CORRIGER"),
    ("petite modif svp", "CORRIGER"),
    ("à revoir", "CORRIGER"),
    ("???", "INCONNU"),
    ("peut-être", "INCONNU"),
    (None, None),
    ("   ", None),
])
def test_lire_decision(cellule, attendu):
    assert review.lire_decision(cellule) == attendu


def test_normaliser_texte_gomme_les_fins_de_ligne_windows_et_les_espaces():
    assert review.normaliser_texte("Bonjour  \r\nà tous   \n") == "Bonjour\nà tous"
    assert review.normaliser_texte(None) == ""


def test_empreinte_change_avec_la_decision_et_la_legende():
    a = review.empreinte("OK", "texte", "18:00")
    assert len(a) == 16
    assert a == review.empreinte("OK", "texte  \r\n", "18:00")
    assert a != review.empreinte("NON", "texte", "18:00")
    assert a != review.empreinte("OK", "autre", "18:00")
    assert a != review.empreinte("OK", "texte", None)


# ---------------------------------------------------------------- rapprochement

def _post(pid, dt, draft=True, **extra):
    return {"id": pid, "draft": draft, "publicationDate": {"dateTime": dt}, **extra}


def test_retrouver_post_par_id_puis_par_creneau():
    original = {"post_id": 1, "date_time": "2026-09-07T18:00:00"}
    posts = [_post(1, "2026-09-07T18:00:00"), _post(2, "2026-09-08T18:00:00")]
    assert review.retrouver_post(posts, original)["id"] == 1

    # L'id a changé (brouillon recréé) : un seul brouillon sur le créneau suffit.
    posts = [_post(9, "2026-09-07T18:00:00")]
    assert review.retrouver_post(posts, original)["id"] == 9

    # Deux brouillons sur le même créneau : on ne devine pas.
    posts = [_post(9, "2026-09-07T18:00:00"), _post(10, "2026-09-07T18:00:00")]
    assert review.retrouver_post(posts, original) is None

    # Le post du créneau n'est plus un brouillon : il ne se repêche pas.
    posts = [_post(9, "2026-09-07T18:00:00", draft=False)]
    assert review.retrouver_post(posts, original) is None


def test_corps_de_sort_du_brouillon_et_garde_le_type():
    post = {"id": 1, "media": ["https://x/1.jpg"], "instagramData": {"type": "REEL"},
            "providers": [{"network": "instagram", "status": "PENDING"}, {"status": "?"}],
            "publicationDate": {"dateTime": "2026-09-07T18:00:00", "timezone": "Europe/Paris"},
            "firstCommentText": "premier commentaire"}
    corps = review.corps_de(post, "nouveau texte", "2026-09-07T19:00:00", True)
    assert corps["draft"] is False and corps["autoPublish"] is True
    assert corps["text"] == "nouveau texte"
    assert corps["publicationDate"] == {"dateTime": "2026-09-07T19:00:00", "timezone": "Europe/Paris"}
    assert corps["instagramData"] == {"type": "REEL"}
    assert corps["providers"] == [{"network": "instagram"}]
    assert corps["media"] == ["https://x/1.jpg"]
    assert corps["firstCommentText"] == "premier commentaire"


# ---------------------------------------------------------------- aller-retour Excel

COMPTE = {"slug": "test-client", "nom": "Client Test", "prefixe": "TC", "blog_id": 42}


def _brouillons():
    return [
        {"id": 101, "draft": True, "text": "Lundi : le post du lundi",
         "media": ["https://cdn/1.jpg"], "instagramData": {"type": "POST"},
         "providers": [{"network": "instagram"}, {"network": "facebook"}],
         "publicationDate": {"dateTime": "2026-09-07T18:00:00"}},
        {"id": 102, "draft": True, "text": "Story sans média\navec deux lignes",
         "media": [], "instagramData": {"type": "STORY"},
         "providers": [{"network": "instagram"}],
         "publicationDate": {"dateTime": "2026-09-08T10:00:00"}},
        {"id": 103, "draft": True, "text": "Résultat : [VAINQUEUR] a gagné",
         "media": ["https://cdn/a.jpg", "https://cdn/b.jpg", "https://cdn/c.jpg"],
         "providers": [{"network": "instagram"}],
         "publicationDate": {"dateTime": "2026-09-09T18:00:00"}},
    ]


def test_aller_retour_fichier(tmp_path):
    chemin = tmp_path / "test-client-2026-09-07.xlsx"
    review.ecrire_fichier(COMPTE, date(2026, 9, 7), _brouillons(), chemin)
    assert chemin.exists()

    meta, lignes = review.lire_review(chemin)
    assert meta == {"client": "test-client", "semaine": "2026-09-07"}
    assert [l["ref"] for l in lignes] == ["TC-01", "TC-02", "TC-03"]
    for ligne, post in zip(lignes, _brouillons()):
        assert ligne["original"]["post_id"] == post["id"]
        assert ligne["original"]["blog_id"] == 42
        assert ligne["original"]["date_time"] == post["publicationDate"]["dateTime"]
        assert ligne["original"]["texte"] == post["text"]
        assert ligne["legende"] == post["text"]
        assert ligne["decision"] is None and ligne["commentaire"] is None
        assert review.normaliser_heure(ligne["heure"]) == post["publicationDate"]["dateTime"][11:16]


def test_aller_retour_lit_les_octets_comme_le_fichier(tmp_path):
    chemin = tmp_path / "f.xlsx"
    review.ecrire_fichier(COMPTE, date(2026, 9, 7), _brouillons(), chemin)
    meta_f, lignes_f = review.lire_review(chemin)
    meta_b, lignes_b = review.lire_review(chemin.read_bytes())
    assert meta_f == meta_b and lignes_f == lignes_b


def test_fichier_a_les_trois_onglets_et_le_menu_des_decisions(tmp_path):
    import openpyxl
    chemin = tmp_path / "f.xlsx"
    review.ecrire_fichier(COMPTE, date(2026, 9, 7), _brouillons(), chemin)
    wb = openpyxl.load_workbook(chemin)
    assert wb.sheetnames == ["Mode d'emploi", "Semaine du 07-09", "_tech"]
    assert wb["_tech"].sheet_state == "hidden"
    ws = wb["Semaine du 07-09"]
    assert [c.value for c in ws[1]] == review.ENTETES
    validations = ws.data_validations.dataValidation
    assert len(validations) == 1
    for decision in (review.DECISION_OK, review.DECISION_CORRIGER, review.DECISION_NON):
        assert decision in validations[0].formula1
    assert ws["F2"].hyperlink.target == "https://cdn/1.jpg"
    assert ws["F3"].value == "⚠️ visuel à venir"
    assert ws["D4"].value == "Carrousel (3 visuels)"


def test_lire_review_refuse_un_fichier_etranger(tmp_path):
    import openpyxl
    chemin = tmp_path / "etranger.xlsx"
    openpyxl.Workbook().save(chemin)
    with pytest.raises(RuntimeError, match="_tech"):
        review.lire_review(chemin)


# ---------------------------------------------------------------- configuration réelle

def test_config_des_comptes_actifs_est_complete():
    for compte in review.charger_comptes():
        for champ in ("slug", "nom", "prefixe", "blog_id", "drive_dossier"):
            assert compte.get(champ), f"{compte.get('slug')} : {champ} manquant"
        assert review.fichiers_de(compte), f"{compte['slug']} : ni drive_fichier ni drive_fichiers"
        # la semaine en cours reste ouverte jusqu'au vendredi, en plus des semaines d'avance
        avance = int(compte.get("semaines_avance") or 1)
        assert len(review.fichiers_de(compte)) <= avance + 1, \
            f"{compte['slug']} : plus de fichiers ouverts que la semaine en cours + l'avance"


def test_fichiers_de_fusionne_sans_doublon():
    assert review.fichiers_de({"drive_fichier": "a"}) == ["a"]
    assert review.fichiers_de({"drive_fichiers": ["a", "b"], "drive_fichier": "a"}) == ["a", "b"]
    assert review.fichiers_de({"drive_fichiers": ["a", None, "b"]}) == ["a", "b"]
    assert review.fichiers_de({"drive_fichier": None}) == []


def test_semaines_a_venir_selon_l_avance():
    vendredi = date(2026, 10, 9)
    assert review.semaines_a_venir({}, vendredi) == [date(2026, 10, 12)]
    assert review.semaines_a_venir({"semaines_avance": 2}, vendredi) == [date(2026, 10, 12),
                                                                          date(2026, 10, 19)]
    # un lundi vise le lundi d'après, jamais le jour même
    assert review.semaines_a_venir({}, date(2026, 10, 12)) == [date(2026, 10, 19)]


def test_charger_comptes_refuse_un_inconnu():
    with pytest.raises(SystemExit):
        review.charger_comptes("client-qui-n-existe-pas")


# ---------------------------------------------------------------- validation automatique

def _brouillon(quand, draft=True, media=("m",), texte="Légende propre", kind="POST"):
    return {"id": quand, "draft": draft, "media": list(media), "text": texte,
            "publicationDate": {"dateTime": quand},
            "providers": [{"network": "instagram"}], "instagramData": {"type": kind}}


def test_a_valider_auto_trie_et_garde_les_garde_fous():
    maintenant = datetime(2026, 10, 7, 12, 0)
    posts = [
        _brouillon("2026-10-07T12:10:00"),                       # dans la marge de 30 min
        _brouillon("2026-10-07T10:00:00"),                       # passé
        _brouillon("2026-10-08T12:00:00"),                       # ✓
        _brouillon("2026-10-08T18:00:00", draft=False),          # déjà publiable
        _brouillon("2026-10-09T12:00:00", media=()),             # sans média
        _brouillon("2026-10-10T12:00:00", texte="Bravo [VAINQUEUR]"),
        _brouillon("2026-10-11T12:00:00", texte="Un tiret — ici"),
        _brouillon("2026-10-12T10:00:00", texte="Story lun — Belle semaine", kind="STORY"),  # ✓
        _brouillon("2026-10-30T12:00:00"),                       # au-delà de l'horizon
    ]
    a_publier, retenus = review.a_valider_auto(posts, maintenant)
    assert [p["id"] for p in a_publier] == ["2026-10-08T12:00:00", "2026-10-12T10:00:00"]
    assert [r for _, r in retenus] == ["aucun média attaché", "texte à placeholders",
                                       "tiret cadratin dans la légende (règle de rédaction)"]


def test_comptes_auto_hors_du_circuit_des_tableaux():
    auto = {c["slug"] for c in review.comptes_auto()}
    assert auto == {"otea-production", "otea-filmakers"}
    assert not auto & {c["slug"] for c in review.charger_comptes()}
